"""Knowledge spaces, sources, ingestion state and grounded retrieval."""
import io
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone
from xml.etree import ElementTree

from ..database import get_conn

MAX_BYTES = 50 * 1024 * 1024
MAX_TEXT = 200_000


def now():
    return datetime.now(timezone.utc).isoformat()


def init_knowledge():
    with get_conn() as conn:
        conn.executescript('''
        CREATE TABLE IF NOT EXISTS knowledge_spaces (
          id TEXT PRIMARY KEY, user_id TEXT NOT NULL, name TEXT NOT NULL,
          description TEXT, visibility TEXT NOT NULL DEFAULT 'private', created_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS knowledge_space_owner ON knowledge_spaces(user_id);
        CREATE TABLE IF NOT EXISTS knowledge_space_members (
          id TEXT PRIMARY KEY, space_id TEXT NOT NULL, user_id TEXT NOT NULL,
          member_name TEXT NOT NULL, role TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS knowledge_sources (
          id TEXT PRIMARY KEY, user_id TEXT NOT NULL, space_id TEXT NOT NULL,
          source_type TEXT NOT NULL, name TEXT NOT NULL, uri TEXT,
          status TEXT NOT NULL, last_synced_at TEXT, document_count INTEGER DEFAULT 0,
          error_message TEXT, created_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS knowledge_source_owner ON knowledge_sources(user_id);
        CREATE TABLE IF NOT EXISTS knowledge_sync_jobs (
          id TEXT PRIMARY KEY, user_id TEXT NOT NULL, source_id TEXT NOT NULL,
          status TEXT NOT NULL, stage TEXT NOT NULL, progress INTEGER DEFAULT 0,
          message TEXT, created_at TEXT NOT NULL, finished_at TEXT);
        CREATE TABLE IF NOT EXISTS knowledge_documents (
          id TEXT PRIMARY KEY, user_id TEXT NOT NULL, title TEXT NOT NULL,
          source TEXT NOT NULL, created_at TEXT NOT NULL, analysis TEXT,
          space_id TEXT, source_id TEXT, doc_type TEXT DEFAULT 'document',
          status TEXT DEFAULT 'ready');
        CREATE INDEX IF NOT EXISTS knowledge_owner ON knowledge_documents(user_id);
        CREATE TABLE IF NOT EXISTS knowledge_chunks (
          id TEXT PRIMARY KEY, document_id TEXT NOT NULL, position INTEGER NOT NULL,
          text TEXT NOT NULL, locator TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS knowledge_chunk_doc ON knowledge_chunks(document_id);
        ''')
        cols = {r['name'] for r in conn.execute('PRAGMA table_info(knowledge_documents)')}
        for name, definition in (
            ('space_id', 'TEXT'), ('source_id', 'TEXT'),
            ('doc_type', "TEXT DEFAULT 'document'"), ('status', "TEXT DEFAULT 'ready'"),
        ):
            if name not in cols:
                conn.execute(f'ALTER TABLE knowledge_documents ADD COLUMN {name} {definition}')


def ensure_default_space(uid):
    with get_conn() as conn:
        row = conn.execute(
            'SELECT * FROM knowledge_spaces WHERE user_id=? ORDER BY created_at LIMIT 1', (uid,)
        ).fetchone()
        if row:
            return dict(row)
        space = {
            'id': uuid.uuid4().hex, 'user_id': uid, 'name': '清华优优 AI Wiki',
            'description': '脚本、直播复盘、经营报表与竞品情报的统一知识空间',
            'visibility': 'private', 'created_at': now(),
        }
        conn.execute(
            'INSERT INTO knowledge_spaces VALUES (:id,:user_id,:name,:description,:visibility,:created_at)', space
        )
        conn.execute(
            'INSERT INTO knowledge_space_members VALUES (?,?,?,?,?,?)',
            (uuid.uuid4().hex, space['id'], uid, '空间所有者', 'owner', now()),
        )
        return space


def create_space(uid, name, description='', visibility='private'):
    space = {
        'id': uuid.uuid4().hex, 'user_id': uid, 'name': name[:100],
        'description': description[:500], 'visibility': visibility, 'created_at': now(),
    }
    with get_conn() as conn:
        conn.execute(
            'INSERT INTO knowledge_spaces VALUES (:id,:user_id,:name,:description,:visibility,:created_at)', space
        )
        conn.execute(
            'INSERT INTO knowledge_space_members VALUES (?,?,?,?,?,?)',
            (uuid.uuid4().hex, space['id'], uid, '空间所有者', 'owner', now()),
        )
    return space


def list_spaces(uid):
    ensure_default_space(uid)
    with get_conn() as conn:
        spaces = [dict(r) for r in conn.execute(
            '''SELECT s.*,COUNT(DISTINCT d.id) AS document_count,COUNT(DISTINCT m.id) AS member_count
               FROM knowledge_spaces s
               LEFT JOIN knowledge_documents d ON d.space_id=s.id
               LEFT JOIN knowledge_space_members m ON m.space_id=s.id
               WHERE s.user_id=? GROUP BY s.id ORDER BY s.created_at''', (uid,)
        )]
        for space in spaces:
            space['members'] = [dict(r) for r in conn.execute(
                'SELECT id,member_name,role,created_at FROM knowledge_space_members WHERE space_id=?',
                (space['id'],),
            )]
        return spaces


def add_space_member(uid, space_id, member_name, role):
    with get_conn() as conn:
        if not conn.execute('SELECT 1 FROM knowledge_spaces WHERE id=? AND user_id=?', (space_id, uid)).fetchone():
            return None
        member = {'id': uuid.uuid4().hex, 'space_id': space_id, 'user_id': uid,
                  'member_name': member_name[:100], 'role': role, 'created_at': now()}
        conn.execute(
            'INSERT INTO knowledge_space_members VALUES (:id,:space_id,:user_id,:member_name,:role,:created_at)', member
        )
        return member


def create_source(uid, source_type, name, uri='', space_id=None, status='syncing'):
    space_id = space_id or ensure_default_space(uid)['id']
    source = {
        'id': uuid.uuid4().hex, 'user_id': uid, 'space_id': space_id,
        'source_type': source_type, 'name': name[:200], 'uri': uri[:2000],
        'status': status, 'last_synced_at': None, 'document_count': 0,
        'error_message': None, 'created_at': now(),
    }
    with get_conn() as conn:
        conn.execute('''INSERT INTO knowledge_sources
          (id,user_id,space_id,source_type,name,uri,status,last_synced_at,document_count,error_message,created_at)
          VALUES (:id,:user_id,:space_id,:source_type,:name,:uri,:status,:last_synced_at,:document_count,:error_message,:created_at)''', source)
    return source


def start_sync_job(uid, source_id, message='等待解析'):
    job = {'id': uuid.uuid4().hex, 'user_id': uid, 'source_id': source_id,
           'status': 'running', 'stage': 'parse', 'progress': 10,
           'message': message, 'created_at': now(), 'finished_at': None}
    with get_conn() as conn:
        conn.execute('''INSERT INTO knowledge_sync_jobs
          (id,user_id,source_id,status,stage,progress,message,created_at,finished_at)
          VALUES (:id,:user_id,:source_id,:status,:stage,:progress,:message,:created_at,:finished_at)''', job)
    return job


def finish_sync_job(uid, source_id, ok=True, message='同步完成'):
    status = 'ready' if ok else 'failed'
    with get_conn() as conn:
        conn.execute(
            '''UPDATE knowledge_sync_jobs SET status=?,stage=?,progress=?,message=?,finished_at=?
               WHERE id=(SELECT id FROM knowledge_sync_jobs WHERE user_id=? AND source_id=? ORDER BY created_at DESC LIMIT 1)''',
            (status, 'index' if ok else 'error', 100 if ok else 0, message, now(), uid, source_id),
        )
        conn.execute(
            '''UPDATE knowledge_sources SET status=?,last_synced_at=?,error_message=?,
               document_count=(SELECT COUNT(*) FROM knowledge_documents WHERE source_id=?)
               WHERE id=? AND user_id=?''',
            (status, now() if ok else None, None if ok else message, source_id, source_id, uid),
        )


def list_sources(uid):
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            '''SELECT s.*,sp.name AS space_name FROM knowledge_sources s
               LEFT JOIN knowledge_spaces sp ON sp.id=s.space_id
               WHERE s.user_id=? ORDER BY s.created_at DESC''', (uid,)
        )]


def list_sync_jobs(uid, limit=50):
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            '''SELECT j.*,s.name AS source_name,s.source_type FROM knowledge_sync_jobs j
               LEFT JOIN knowledge_sources s ON s.id=j.source_id
               WHERE j.user_id=? ORDER BY j.created_at DESC LIMIT ?''', (uid, limit)
        )]


def parse_file(name, data):
    suffix = name.lower().rsplit('.', 1)[-1]
    if len(data) > MAX_BYTES:
        raise ValueError('文件不能超过 50 MB')
    if suffix in ('docx', 'xlsx', 'pptx'):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError('文件解压后过大，请拆分上传')
    if suffix == 'docx':
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = ElementTree.fromstring(archive.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        sections = [('段落 %s' % (i + 1), ''.join(t.text or '' for t in p.findall('.//w:t', ns))) for i, p in enumerate(root.findall('.//w:p', ns))]
    elif suffix == 'pptx':
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = sorted([n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)], key=lambda n:int(re.search(r'slide(\d+)',n).group(1)))
            sections = []
            for index, name in enumerate(names, 1):
                root = ElementTree.fromstring(archive.read(name))
                paragraphs = [''.join(t.text or '' for t in p.iter() if t.tag.rsplit('}',1)[-1]=='t') for p in root.iter() if p.tag.rsplit('}',1)[-1]=='p']
                sections.append((f'第 {index} 页', '\n'.join(paragraphs)))
    elif suffix == 'pdf':
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError('PDF 解析组件未安装，请执行 pip install -r requirements.txt') from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError('加密 PDF 无法解析，请导出未加密版本后上传')
            if len(reader.pages) > 500:
                raise ValueError('PDF 页数超过 500 页，请拆分上传')
            sections = [(f'第 {i + 1} 页', page.extract_text() or '') for i, page in enumerate(reader.pages)]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('PDF 解析失败，请确认文件未损坏') from exc
    elif suffix == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections, total = [], 0
        try:
            for sheet in book:
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    text = ' | '.join(str(v) if v is not None else '' for v in row)
                    sections.append((f'{sheet.title} · 第 {i} 行', text))
                    total += len(text)
                    if total > MAX_TEXT or len(sections) > 20000:
                        raise ValueError('表格过大，请拆分上传')
        finally:
            book.close()
    elif suffix in ('txt', 'md', 'csv', 'tsv', 'srt', 'vtt'):
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('gb18030')
        sections = [(f'第 {i} 行', line) for i, line in enumerate(text.splitlines(), 1)]
    else:
        raise ValueError('支持 PDF、DOCX、XLSX、PPTX、TXT、MD、CSV、TSV、SRT、VTT；旧版 Word/Excel 请另存为 DOCX/XLSX')
    if sum(len(text) for _, text in sections) > MAX_TEXT:
        raise ValueError('正文不能超过 20 万字，请拆分上传')
    sections = [(loc, text.strip()) for loc, text in sections if text.strip()]
    if not sections:
        raise ValueError('没有提取到文本；扫描版 PDF 和图片需先 OCR 转成文字')
    return sections


def split_sections(sections, size=800, overlap=100):
    chunks = []
    for locator, text in sections:
        for start in range(0, len(text), size - overlap):
            chunks.append({'text': text[start:start + size], 'locator': locator})
            if start + size >= len(text):
                break
    merged = []
    for chunk in chunks:
        if merged and len(merged[-1]['text']) + len(chunk['text']) + 1 <= size:
            merged[-1]['text'] += '\n' + chunk['text']
            merged[-1]['locator'] = merged[-1]['locator'].split(' → ')[0] + ' → ' + chunk['locator']
        else:
            merged.append(dict(chunk))
    return merged


def save_document(uid, title, source, sections, space_id=None, source_id=None, doc_type='document'):
    chunks = split_sections(sections)
    if not chunks:
        raise ValueError('没有可入库内容')
    space_id = space_id or ensure_default_space(uid)['id']
    doc_id = uuid.uuid4().hex
    with get_conn() as conn:
        conn.execute(
            '''INSERT INTO knowledge_documents
               (id,user_id,title,source,created_at,analysis,space_id,source_id,doc_type,status)
               VALUES (?,?,?,?,?,NULL,?,?,?,'ready')''',
            (doc_id, uid, title[:200], source, now(), space_id, source_id, doc_type),
        )
        conn.executemany(
            'INSERT INTO knowledge_chunks VALUES (?,?,?,?,?)',
            [(uuid.uuid4().hex, doc_id, i + 1, c['text'], c['locator']) for i, c in enumerate(chunks)],
        )
    return get_document(uid, doc_id)


def list_documents(uid):
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            '''SELECT d.id,d.title,d.source,d.created_at,d.analysis IS NOT NULL AS analyzed,
                      d.space_id,d.source_id,d.doc_type,d.status,sp.name AS space_name,
                      COUNT(c.id) AS chunk_count
               FROM knowledge_documents d
               LEFT JOIN knowledge_chunks c ON c.document_id=d.id
               LEFT JOIN knowledge_spaces sp ON sp.id=d.space_id
               WHERE d.user_id=? GROUP BY d.id ORDER BY d.created_at DESC''', (uid,)
        )]


def get_document(uid, doc_id):
    with get_conn() as conn:
        row = conn.execute(
            '''SELECT d.*,sp.name AS space_name,s.name AS source_name
               FROM knowledge_documents d
               LEFT JOIN knowledge_spaces sp ON sp.id=d.space_id
               LEFT JOIN knowledge_sources s ON s.id=d.source_id
               WHERE d.id=? AND d.user_id=?''', (doc_id, uid)
        ).fetchone()
        if not row:
            return None
        result = dict(row)
        result.pop('user_id')
        result['analysis'] = json.loads(result['analysis']) if result['analysis'] else None
        result['chunks'] = [dict(r) for r in conn.execute(
            'SELECT id,position,text,locator FROM knowledge_chunks WHERE document_id=? ORDER BY position', (doc_id,)
        )]
        return result


def terms(text):
    lowered = text.lower()
    words = re.findall(r'[a-z0-9_]+', lowered)
    for run in re.findall(r'[\u4e00-\u9fff]+', lowered):
        words.extend(run[i:i + 2] for i in range(max(1, len(run) - 1)))
    return set(words)


def retrieve(uid, question, limit=8, space_id=None, document_ids=None):
    query = terms(question)
    if not query and not document_ids:
        return []
    sql = '''SELECT c.*,d.title,d.space_id,d.source_id,d.doc_type
             FROM knowledge_chunks c JOIN knowledge_documents d ON d.id=c.document_id
             WHERE d.user_id=?'''
    args = [uid]
    if space_id:
        sql += ' AND d.space_id=?'
        args.append(space_id)
    # Cue 的 @文档是一个明确的检索边界：只从用户选中的资料中取证，
    # 避免同名脚本、旧报表等其他资料混入回答。
    if document_ids:
        document_ids = [str(doc_id) for doc_id in document_ids[:20] if str(doc_id).strip()]
        if document_ids:
            sql += ' AND d.id IN (' + ','.join('?' for _ in document_ids) + ')'
            args.extend(document_ids)
    with get_conn() as conn:
        rows = conn.execute(sql, args).fetchall()
    scored = []
    for row in rows:
        text_terms = terms(row['text'] + ' ' + row['title'])
        score = len(query & text_terms) / max(1, len(query))
        if score > 0 or document_ids:
            scored.append((score, dict(row)))
    scored.sort(key=lambda pair: pair[0], reverse=True)
    return [dict(row, score=round(score, 3)) for score, row in scored[:limit]]


def knowledge_graph(uid):
    nodes, edges, seen_nodes, seen_edges = [], [], set(), set()
    with get_conn() as conn:
        rows = conn.execute(
            'SELECT id,title,analysis FROM knowledge_documents WHERE user_id=? AND analysis IS NOT NULL', (uid,)
        ).fetchall()
    for row in rows:
        try:
            analysis = json.loads(row['analysis'])
        except (TypeError, ValueError):
            continue
        for node in analysis.get('nodes', []):
            label = str(node.get('label', '')).strip()
            if not label:
                continue
            key = label.lower()
            if key not in seen_nodes:
                nodes.append({'id': key, 'label': label, 'kind': node.get('kind', '概念')})
                seen_nodes.add(key)
        local = {str(n.get('id')): str(n.get('label', '')).strip().lower() for n in analysis.get('nodes', [])}
        with get_conn() as conn:
            chunks = {c['position']: c['text'] for c in conn.execute('SELECT position,text FROM knowledge_chunks WHERE document_id=?', (row['id'],))}
        for edge in analysis.get('edges', []):
            quote = str(edge.get('quote', '')).strip()
            if not quote or quote not in chunks.get(edge.get('chunk'), ''):
                continue
            source, target = local.get(str(edge.get('source'))), local.get(str(edge.get('target')))
            if source in seen_nodes and target in seen_nodes:
                key = (source, target, edge.get('relation'))
                if key not in seen_edges:
                    edges.append({'source': source, 'target': target, 'relation': edge.get('relation', '关联'),
                                  'document_id': row['id'], 'document_title': row['title'],
                                  'quote': quote, 'chunk': edge.get('chunk'), 'evidence_status': 'quote_verified'})
                    seen_edges.add(key)
    visible = {n['id'] for n in nodes[:80]}
    return {'nodes': nodes[:80], 'edges': [e for e in edges if e['source'] in visible and e['target'] in visible][:120], 'document_count': len(rows)}
